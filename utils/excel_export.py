import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from utils.data_processing import MONTHS_SPANISH, MONTH_ABBR

DARK='1B365D'; BLUE='D9EAF7'; LIGHT='F3F6FA'; WHITE='FFFFFF'; GREEN='E2F0D9'
THIN=Side(style='thin', color='D9E1F2')


def _style(cell, fill=None, bold=False, color='000000', align='center'):
    cell.font=Font(name='Calibri',size=10,bold=bold,color=color)
    if fill: cell.fill=PatternFill('solid',fgColor=fill)
    cell.alignment=Alignment(horizontal=align,vertical='center',wrap_text=True)
    cell.border=Border(left=THIN,right=THIN,top=THIN,bottom=THIN)


def generate_excel_download(pivot_df, kpis, filters_info):
    wb=Workbook(); ws=wb.active; ws.title='Detalle'; ws.sheet_view.showGridLines=False
    ws['A1']='SEGUIMIENTO DE METAS Y EJECUCIÓN - RED CITE PÚBLICA'; ws['A1'].font=Font(size=15,bold=True,color=DARK)
    ws.merge_cells('A1:AC1')
    ws['A2']=f"Año: {filters_info.get('Año')} | Indicador: {filters_info.get('Indicador')} | Categoría: {filters_info.get('Categoría presupuestal')} | CITE/UT: {filters_info.get('CITE/UT')}"
    ws.merge_cells('A2:AC2')
    ws['A3']=f"Periodo: {filters_info.get('Mes de inicio')} – {filters_info.get('Mes de fin')} | Tipo servicio: {filters_info.get('Tipo servicio')} | Tarea: {filters_info.get('Tarea')} | Complejidad: {filters_info.get('Complejidad')}"
    ws.merge_cells('A3:AC3')
    r1,r2=5,6; ws.merge_cells(start_row=r1,start_column=1,end_row=r2,end_column=1); ws.cell(r1,1,'Detalle')
    c=2
    for m in MONTHS_SPANISH:
        ws.merge_cells(start_row=r1,start_column=c,end_row=r1,end_column=c+1); ws.cell(r1,c,MONTH_ABBR[m].upper())
        ws.cell(r2,c,'Meta'); ws.cell(r2,c+1,'Ejecución'); c+=2
    for title in ['TOTAL PERIODO','TOTAL ANUAL']:
        ws.merge_cells(start_row=r1,start_column=c,end_row=r1,end_column=c+2); ws.cell(r1,c,title)
        ws.cell(r2,c,'Meta'); ws.cell(r2,c+1,'Ejecución'); ws.cell(r2,c+2,'% Avance'); c+=3
    max_col=c-1
    for row in ws.iter_rows(min_row=r1,max_row=r2,min_col=1,max_col=max_col):
        for cell in row: _style(cell,DARK,True,WHITE)
    data_start=7
    if pivot_df is not None and not pivot_df.empty:
        for i,(_,row) in enumerate(pivot_df.iterrows(),start=data_start):
            ws.cell(i,1,row['Detalle']); _style(ws.cell(i,1), LIGHT if row['Detalle']=='TOTAL' else None, row['Detalle']=='TOTAL', align='left')
            c=2
            for m in MONTHS_SPANISH:
                for key in [f'{m}_Meta',f'{m}_Ejec']:
                    cell=ws.cell(i,c,float(row.get(key,0))); cell.number_format='#,##0'; _style(cell, LIGHT if row['Detalle']=='TOTAL' else None,row['Detalle']=='TOTAL'); c+=1
            for keys in [('Meta_Periodo','Ejec_Periodo','Avance_Periodo'),('Meta_Anual','Ejec_Anual','Avance_Anual')]:
                for j,key in enumerate(keys):
                    val=float(row.get(key,0)); cell=ws.cell(i,c,val/100 if 'Avance' in key else val); cell.number_format='0.0%' if 'Avance' in key else '#,##0'; _style(cell,GREEN if 'Avance' in key else (LIGHT if row['Detalle']=='TOTAL' else None),row['Detalle']=='TOTAL'); c+=1
    ws.freeze_panes='B7'; ws.auto_filter.ref=f'A6:{get_column_letter(max_col)}{max(data_start, data_start + (len(pivot_df) if pivot_df is not None else 0))}'
    ws.column_dimensions['A'].width=34
    for col in range(2,max_col+1): ws.column_dimensions[get_column_letter(col)].width=12

    wr=wb.create_sheet('Resumen'); wr.sheet_view.showGridLines=False
    wr['A1']='RESUMEN DE CONSULTA'; wr['A1'].font=Font(size=15,bold=True,color=DARK); wr.merge_cells('A1:C1')
    info=[('Meta del periodo',kpis['meta_periodo'],'#,##0'),('Ejecutado del periodo',kpis['ejec_periodo'],'#,##0'),('Avance del periodo',kpis['avance_periodo']/100,'0.0%'),('Brecha del periodo',kpis['brecha_periodo'],'#,##0'),('Meta anual',kpis['meta_anual'],'#,##0'),('Ejecutado anual',kpis['ejec_anual'],'#,##0'),('Avance anual',kpis['avance_anual']/100,'0.0%')]
    for i,(label,val,fmt) in enumerate(info,start=3):
        wr.cell(i,1,label); wr.cell(i,2,val); wr.cell(i,2).number_format=fmt; _style(wr.cell(i,1),LIGHT,True,align='left'); _style(wr.cell(i,2))
    wr.column_dimensions['A'].width=28; wr.column_dimensions['B'].width=18
    out=io.BytesIO(); wb.save(out); out.seek(0); return out.getvalue()
